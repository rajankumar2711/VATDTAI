"""
Helpers for the e-Invoice validation-agent UI E2E flow.

Track B of the e-invoice agent testing effort: drives the VATdtai app end to end
(Data Ingestion -> e-Invoice Management -> agent -> Review Required) and compares the
agent's Review Required results against the master expected-results workbook.

File discovery + offline comparison live here so the slow, QA-mutating live capture
runs once and the comparison against the golden workbook can be re-run cheaply.
"""
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, List

PROJECT_ROOT = Path(__file__).resolve().parent.parent
# Quarantined in Phase 3: test-data + expected-results now live next to this module
# under agent-quarantine/einvoice-agent-testing/.
EAT_DIR = Path(__file__).resolve().parent / "einvoice-agent-testing"
TEST_DATA_DIR = EAT_DIR / "test-data"
EXPECTED_WORKBOOK = EAT_DIR / "expected results" / "EInvoice_Agent_Test_Master_60_Samples_v3.xlsx"
RESULTS_DIR = EAT_DIR / "results"
RAW_DIR = RESULTS_DIR / "raw"
REPORTS_DIR = RESULTS_DIR / "reports"

# Source-system folder -> Data Ingestion dropdown option label (exact text in the
# 'Select Source System' dropdown: "SAP", "Custom ERP", "MS D365").
SYSTEM_TO_SOURCE = {"SAP": "SAP", "CustomERP": "Custom ERP"}

COUNTRIES = ("Belgium", "France", "Poland")


# ==========================================
# FILE DISCOVERY
# ==========================================
def discover_country_files(country: str) -> List[Dict]:
    """Return every test-data XML for a country with its Data Ingestion source system."""
    base = TEST_DATA_DIR / country
    files: List[Dict] = []
    for system_folder in ("SAP", "CustomERP"):
        folder = base / system_folder
        if not folder.exists():
            continue
        for xml in sorted(folder.glob("*.xml")):
            files.append({
                "path": str(xml),
                "file_name": xml.name,
                "system_folder": system_folder,
                "source_system": SYSTEM_TO_SOURCE[system_folder],
                # Filename convention: <System>_<CC>_<NN>_<INVOICE-NUMBER>.xml
                "invoice_number": xml.stem.split("_")[-1],
            })
    return files


def new_run_token() -> str:
    """Short run-unique token used to mint fresh invoice numbers."""
    return datetime.now().strftime("R%m%d%H%M%S")


def make_fresh_files(files: List[Dict], run_token: str):
    """
    Write a run-scoped copy of each test-data XML with its invoice number replaced by a
    fresh, run-unique value (``<original>-<run_token>``) so the agent RE-EVALUATES every
    scenario cleanly instead of reading stale statuses / duplicate rows left by earlier
    runs that reused the same numbers.

    Files that intentionally SHARE an invoice number (duplicate-detection scenarios)
    keep sharing the SAME fresh number, so the duplicate condition is preserved.

    Returns (fresh_files, alias_map) where alias_map maps fresh_number -> original_number
    so the offline comparison can still resolve expected findings by the original id.
    """
    out_dir = RESULTS_DIR / "_fresh_xml" / run_token
    fresh_files: List[Dict] = []
    alias: Dict[str, str] = {}
    for f in files:
        orig = f["invoice_number"]
        fresh = f"{orig}-{run_token}"
        sub = out_dir / f["system_folder"]
        sub.mkdir(parents=True, exist_ok=True)
        src = Path(f["path"])
        content = src.read_text(encoding="utf-8").replace(orig, fresh)
        dst = sub / src.name
        dst.write_text(content, encoding="utf-8")
        nf = dict(f)
        nf["path"] = str(dst)
        nf["invoice_number"] = fresh
        nf["original_number"] = orig
        fresh_files.append(nf)
        alias[fresh] = orig
    return fresh_files, alias


# ==========================================
# EXPECTED (GOLDEN) DATA
# ==========================================
def load_expected_lookup():
    """Reuse the Phase 1 golden_loader to build invoice_id -> expected records."""
    if str(EAT_DIR) not in sys.path:
        sys.path.insert(0, str(EAT_DIR))
    from src.loaders.golden_loader import load_golden_data  # noqa: E402

    golden = load_golden_data()
    lookup: Dict[str, List[Dict]] = {}
    for sample in golden.samples.values():
        lookup.setdefault(sample.invoice_id, []).append({
            "sample_id": sample.sample_id,
            "country": sample.country,
            "system": sample.system,
            "scenario": sample.scenario,
            "expected_status": sample.expected_status,
            "expected_review": "review" in (sample.expected_status or "").lower(),
            "expected_findings": sample.expected_findings,
            "expected_finding_count": sample.expected_finding_count,
        })
    return lookup, golden


def expected_invoice_ids(country: str) -> set:
    """Golden invoice ids for a country (used to scope the Review Required drill-in)."""
    lookup, _ = load_expected_lookup()
    return {
        inv for inv, recs in lookup.items()
        if any(_country_matches(r["country"], country) for r in recs)
    }


# Observed anomaly-popup rule-title keyword -> business rule / agent name (as used
# in the master workbook's Expected_Findings sheet). Extend as new popup rule titles
# are seen; unmatched titles are reported as 'UNMAPPED' so the map can be completed.
OBSERVED_RULE_TO_AGENT = [
    ("numbering integrity", "Number Sequence"),
    ("number sequence", "Number Sequence"),
    ("mathematical integrity", "Totals"),
    ("totals", "Totals"),
    ("completeness", "XML Readiness"),
    ("xml readiness", "XML Readiness"),
    ("structural", "XML Readiness"),
    ("hierarchy", "XML Readiness"),
    ("character safety", "Text Safety"),
    ("free-text", "Text Safety"),
    ("free text", "Text Safety"),
    ("unicode", "Text Safety"),
    ("text safety", "Text Safety"),
    ("invoice type", "Invoice Type"),
    ("document type", "Invoice Type"),
    ("credit note", "Invoice Type"),
    ("customer type", "Customer Type"),
    ("customer classification", "Customer Type"),
    ("party", "Customer Type"),
]


def classify_observed_agent(rule_title: str) -> str:
    """Map an observed anomaly popup rule title to a workbook agent/business rule."""
    t = (rule_title or "").strip().lower()
    for kw, agent in OBSERVED_RULE_TO_AGENT:
        if kw in t:
            return agent
    return "UNMAPPED"


def _expected_findings_by_invoice(golden):
    """invoice_id -> {'agents': set, 'issue_types': set} from Expected_Findings."""
    out: Dict[str, Dict[str, set]] = {}
    for ef in golden.expected_findings:
        sample = golden.samples.get(ef.sample_id)
        if not sample or not sample.invoice_id:
            continue
        rec = out.setdefault(sample.invoice_id, {"agents": set(), "issue_types": set()})
        if ef.agent:
            rec["agents"].add(ef.agent)
        if ef.issue_type:
            rec["issue_types"].add(ef.issue_type)
    return out


def _norm_status(status: str) -> str:
    text = (status or "").strip().lower()
    if "review required" in text:
        return "review_required"
    if "reviewed" in text:
        return "reviewed"
    if "ready" in text:
        return "ready"
    if not text:
        return "unknown"
    return text


# Outbound statuses that mean the agent found NO anomalies (a "good"/clean invoice).
_GOOD_STATUSES = ("ready", "review not needed")


def _is_good_status(status: str) -> bool:
    text = (status or "").strip().lower()
    return any(g in text for g in _GOOD_STATUSES)


def _is_review_required_status(status: str) -> bool:
    return "review required" in (status or "").strip().lower()


def _country_matches(record_country: str, country: str) -> bool:
    return (record_country or "").strip().lower() == country.strip().lower()


# ==========================================
# COMPARISON
# ==========================================
def build_report(
    country: str,
    observations: List[Dict],
    ingestion_results: List[Dict] | None = None,
    alias_map: Dict[str, str] | None = None,
) -> Dict:
    """
    Compare, one invoice at a time, the agent's observed Outbound status (and captured
    anomalies) against the master expected-results workbook.

    observations: one dict per Client Invoice Number processed, e.g.
      {invoice_id, source_system, found (bool), observed_status (raw),
       observed_review (bool), anomaly_count, anomalies:[...]}.
    Returns a structured report dict (also persisted by save_report).
    """
    lookup, golden = load_expected_lookup()
    expected_for_country = {
        inv: recs for inv, recs in lookup.items()
        if any(_country_matches(r["country"], country) for r in recs)
    }
    exp_findings = _expected_findings_by_invoice(golden)
    alias_map = alias_map or {}

    # Group observations per Client e-Invoice Number (one entry per physical Outbound
    # row, ordered by occurrence) so intentional duplicates are paired with their
    # multiple expected samples.
    obs_by_inv: Dict[str, List[Dict]] = {}
    for o in observations:
        inv = (o.get("invoice_id") or "").strip()
        if not inv:
            continue
        obs_by_inv.setdefault(inv, []).append(o)
    for inv in obs_by_inv:
        obs_by_inv[inv].sort(key=lambda x: x.get("occurrence", 1))

    per_invoice = []
    matched = mismatched = not_found = no_expected = 0
    findings_matched = findings_missing_total = findings_extra_total = 0
    passed = failed = 0
    processed_ids = set()

    for inv, occs in obs_by_inv.items():
        processed_ids.add(inv)
        # In fresh-number mode the observed id carries a run suffix; resolve expected
        # data by the original invoice number via the alias map.
        key = alias_map.get(inv, inv)
        recs = expected_for_country.get(key) or lookup.get(key) or []
        expected_agents = set(exp_findings.get(key, {}).get("agents", set()))
        # Pair each expected sample with an occurrence; evaluate every physical row and
        # every expected sample (extras on either side handled gracefully).
        pair_count = max(len(recs), len(occs), 1)
        for i in range(pair_count):
            rec = recs[i] if i < len(recs) else (recs[0] if recs else None)
            obs = occs[i] if i < len(occs) else None
            expected_review = rec["expected_review"] if rec else None

            found = bool(obs and obs.get("found", obs.get("observed_status") is not None))
            observed_status = obs.get("observed_status") if obs else None
            observed_review = bool(obs.get("observed_review")) if (obs and found) else None
            anomalies = obs.get("anomalies", []) if obs else []

            if rec is None:
                outcome = "NO_EXPECTED"
                no_expected += 1
            elif not found:
                outcome = "NOT_FOUND"
                not_found += 1
            elif observed_review == expected_review:
                outcome = "MATCH"
                matched += 1
            else:
                outcome = "MISMATCH"
                mismatched += 1

            # ---- Findings-level (agent/business-rule) comparison ----
            observed_agents = {classify_observed_agent(a.get("rule")) for a in anomalies}
            observed_agents_mapped = {a for a in observed_agents if a != "UNMAPPED"}
            f_missing = sorted(expected_agents - observed_agents_mapped)
            f_extra = sorted(observed_agents_mapped - expected_agents)
            if found:
                # "all expected present" (subset) counts as a findings match; extras
                # are recorded but do not by themselves fail the invoice.
                findings_outcome = "FINDINGS_MATCH" if not f_missing else "FINDINGS_MISMATCH"
                if not f_missing:
                    findings_matched += 1
                findings_missing_total += len(f_missing)
                findings_extra_total += len(f_extra)
            else:
                findings_outcome = "NOT_FOUND"

            # ---- Authoritative PASS/FAIL rule ----
            #  - Ready / Review not needed -> PASS (agent found no anomalies)
            #  - Review Required           -> PASS if ALL expected findings are present
            #                                 (subset match); extras are noted, not failed
            #  - Not found / any other     -> FAIL (NOT_FOUND labelled separately)
            if not found:
                test_result = "NOT_FOUND"
                failed += 1
                test_reason = "invoice not found in Outbound grid"
            elif _is_good_status(observed_status):
                # Ready is only a PASS when the invoice was EXPECTED to be clean. If the
                # invoice was expected to be flagged (expected findings / review) but the
                # agent reports Ready, that is a missed detection (false negative) = FAIL.
                if expected_review or expected_agents:
                    test_result = "FAIL"
                    failed += 1
                    exp_txt = sorted(expected_agents) if expected_agents else "review required"
                    test_reason = f"status Ready but expected to be flagged ({exp_txt}) - missed detection"
                else:
                    test_result = "PASS"
                    passed += 1
                    test_reason = "status Ready - clean as expected"
            elif _is_review_required_status(observed_status):
                if expected_agents and expected_agents.issubset(observed_agents_mapped):
                    test_result = "PASS"
                    passed += 1
                    test_reason = ("Review Required - all expected findings present"
                                   + (f" (extra: {f_extra})" if f_extra else ""))
                else:
                    test_result = "FAIL"
                    failed += 1
                    if not expected_agents:
                        test_reason = "Review Required but invoice expected clean (no findings)"
                    else:
                        test_reason = f"findings mismatch (missing={f_missing}, extra={f_extra})"
            else:
                test_result = "FAIL"
                failed += 1
                test_reason = f"unexpected status '{observed_status}'"

            per_invoice.append({
                "test_result": test_result,
                "test_reason": test_reason,
                "invoice_id": inv,
                "original_invoice": key,
                "occurrence": (obs.get("occurrence") if obs else i + 1),
                "sample_id": rec["sample_id"] if rec else None,
                "system": rec["system"] if rec else (obs.get("source_system") if obs else None),
                "scenario": rec["scenario"] if rec else None,
                "expected_status": rec["expected_status"] if rec else None,
                "expected_review": expected_review,
                "found_in_outbound": found,
                "observed_status": observed_status,
                "observed_review": observed_review,
                "outcome": outcome,
                "expected_agents": sorted(expected_agents),
                "expected_issue_types": rec["expected_findings"] if rec else None,
                "observed_agents": sorted(observed_agents_mapped),
                "unmapped_observed_rules": sorted(
                    {a.get("rule") for a in anomalies
                     if classify_observed_agent(a.get("rule")) == "UNMAPPED"}
                ),
                "findings_missing": f_missing,
                "findings_extra": f_extra,
                "findings_outcome": findings_outcome,
                "observed_anomaly_count": obs.get("anomaly_count", 0) if obs else 0,
                "observed_anomalies": anomalies,
            })

    processed_orig = {alias_map.get(i, i) for i in processed_ids}
    expected_not_processed = sorted(set(expected_for_country) - processed_orig)

    report = {
        "country": country,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "summary": {
            "expected_invoices_for_country": len(expected_for_country),
            "processed_invoices": len(processed_ids),
            "passed": passed,
            "failed": failed,
            "matched": matched,
            "mismatched": mismatched,
            "not_found": not_found,
            "no_expected_entry": no_expected,
            "expected_not_processed": len(expected_not_processed),
            "findings_match": findings_matched,
            "findings_missing_total": findings_missing_total,
            "findings_extra_total": findings_extra_total,
        },
        "per_invoice": per_invoice,
        "expected_not_processed_ids": expected_not_processed,
        "ingestion_results": ingestion_results or [],
    }
    return report


# ==========================================
# PERSISTENCE
# ==========================================
def _ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def save_raw(country: str, name: str, data) -> Path:
    out_dir = RAW_DIR / f"ui_{country.lower()}"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{name}_{_ts()}.json"
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def save_report(report: Dict) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    country = report.get("country", "unknown").lower()
    path = REPORTS_DIR / f"ui_agent_e2e_{country}_{_ts()}.json"
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


# Single persistent actual-results workbook, aligned 1:1 with the expected workbook,
# so the two can be compared side by side. Updated in place on every run.
ACTUAL_WORKBOOK = EAT_DIR / "Actual result.xlsx"

_ACTUAL_HEADERS = [
    "Sample ID", "Country", "System", "Invoice ID", "Scenario",
    "Expected Status", "Actual Status", "Result Reason",
    "Expected Finding Count", "Actual Finding Count",
    "Expected Agents", "Observed Agents", "Findings Missing", "Findings Extra",
    "Findings Result", "Test Result", "Last Run",
]
_FINDINGS_HEADERS = [
    "Invoice ID", "Country", "Observed Status", "Agent (mapped)",
    "Rule (popup title)", "Field", "Current", "Suggested", "Details",
]


def update_actual_results_workbook(report: Dict) -> Path:
    """Create/update the single 'Actual result.xlsx' aligned with the expected workbook.

    Sheet 'Comparison' has one row per expected sample (seeded from the master
    workbook) with expected vs actual columns side by side; sheet 'Findings' holds
    the verbatim agent findings captured from the Status popup. Rows are updated in
    place so multiple country runs accumulate into one file.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    _, golden = load_expected_lookup()
    exp_findings = _expected_findings_by_invoice(golden)
    # Key results by sample_id so duplicate invoice numbers (multiple samples sharing
    # one number) map to their own workbook rows.
    by_sid = {r["sample_id"]: r for r in report.get("per_invoice", []) if r.get("sample_id")}
    run_ts = report.get("generated_at", datetime.now().isoformat(timespec="seconds"))

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")

    def _exp_agents(inv):
        return ", ".join(sorted(exp_findings.get(inv, {}).get("agents", set())))

    if ACTUAL_WORKBOOK.exists():
        wb = openpyxl.load_workbook(ACTUAL_WORKBOOK)
        ws = wb["Comparison"] if "Comparison" in wb.sheetnames else wb.active
        ws.title = "Comparison"
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comparison"
        ws.append(_ACTUAL_HEADERS)

    row_by_sid = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for sid, s in golden.samples.items():
        if sid not in row_by_sid:
            ws.append([
                sid, s.country, s.system, s.invoice_id, s.scenario,
                s.expected_status, "", "NOT_RUN", s.expected_finding_count, "",
                _exp_agents(s.invoice_id), "", "", "", "NOT_RUN", "NOT_RUN", "",
            ])
            row_by_sid[sid] = ws.max_row

    for sid, s in golden.samples.items():
        res = by_sid.get(sid)
        if not res:
            continue
        r = row_by_sid[sid]
        test_result = res.get("test_result", "FAIL")
        ws.cell(r, 7).value = res.get("observed_status")
        ws.cell(r, 8).value = res.get("test_reason")
        ws.cell(r, 10).value = res.get("observed_anomaly_count", 0)
        ws.cell(r, 11).value = _exp_agents(s.invoice_id)
        ws.cell(r, 12).value = ", ".join(res.get("observed_agents") or [])
        ws.cell(r, 13).value = ", ".join(res.get("findings_missing") or [])
        ws.cell(r, 14).value = ", ".join(res.get("findings_extra") or [])
        ws.cell(r, 15).value = res.get("findings_outcome")
        tr_cell = ws.cell(r, 16)
        tr_cell.value = test_result
        tr_cell.fill = pass_fill if test_result == "PASS" else (
            warn_fill if test_result == "NOT_FOUND" else fail_fill)
        ws.cell(r, 17).value = run_ts

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    widths = [12, 10, 10, 22, 10, 16, 16, 14, 14, 14, 22, 22, 22, 22, 16, 12, 20]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    # ---- Findings sheet (verbatim); keep prior countries, refresh this run's invoices
    existing_rows = []
    if "Findings" in wb.sheetnames:
        fws_old = wb["Findings"]
        existing_rows = [[c.value for c in row] for row in fws_old.iter_rows(min_row=2)]
        wb.remove(fws_old)
    fws = wb.create_sheet("Findings")
    fws.append(_FINDINGS_HEADERS)
    updated_invs = {r.get("invoice_id") for r in report.get("per_invoice", [])}
    for row in existing_rows:
        if row and row[0] not in updated_invs:
            fws.append(row)
    for r in report.get("per_invoice", []):
        for a in r.get("observed_anomalies") or []:
            details = a.get("details")
            fws.append([
                r["invoice_id"], report.get("country"), r.get("observed_status"),
                classify_observed_agent(a.get("rule")), a.get("rule"),
                a.get("field"), a.get("current"), a.get("suggested"),
                " | ".join(details) if isinstance(details, list) else str(details or ""),
            ])
    for cell in fws[1]:
        cell.font = header_font
        cell.fill = header_fill
    fwidths = [22, 10, 16, 16, 42, 22, 20, 24, 60]
    for idx, w in enumerate(fwidths, start=1):
        fws.column_dimensions[get_column_letter(idx)].width = w

    wb.save(ACTUAL_WORKBOOK)
    return ACTUAL_WORKBOOK


def format_summary(report: Dict) -> str:
    s = report["summary"]
    lines = [
        f"e-Invoice Agent UI E2E - {report['country']}",
        f"  expected invoices for country  : {s['expected_invoices_for_country']}",
        f"  processed (this run)           : {s['processed_invoices']}",
        f"  PASS / FAIL                    : {s.get('passed', 0)} / {s.get('failed', 0)}",
        f"  status MATCH / MISMATCH        : {s['matched']} / {s['mismatched']}",
        f"  NOT_FOUND / NO_EXPECTED        : {s['not_found']} / {s['no_expected_entry']}",
        f"  findings MATCH / missing/extra : {s['findings_match']} / "
        f"{s['findings_missing_total']} / {s['findings_extra_total']}",
        f"  expected not processed         : {s['expected_not_processed']}",
    ]
    for r in report.get("per_invoice", []):
        lines.append(
            f"    [{r.get('test_result')}] {r['invoice_id']} "
            f"obs={r['observed_status']} - {r.get('test_reason', '')}"
        )
        lines.append(
            f"        expected_agents={r['expected_agents']} "
            f"observed_agents={r['observed_agents']} "
            f"missing={r['findings_missing']} extra={r['findings_extra']}"
        )
    return "\n".join(lines)
