"""One-off: render a .docx.json spec into a real .docx using python-docx.

Styling matches the Moonshot AI QA reference templates:
- Body font EYInterstate Light, dark-slate color #243447
- Navy headings (H1 #17365D, H2 #1F4E78, H3 #4F81BD)
- Navy table headers (#1F4E78) with white bold text
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justified": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

BODY_FONT = "EYInterstate Light"
BODY_COLOR = "243447"
H1_COLOR = "17365D"
H2_COLOR = "1F4E78"
H3_COLOR = "4F81BD"
HEADER_FILL = "1F4E78"
BRAND = "2E2E9E"           # legacy purple in specs -> remapped to navy
DARK_FILLS = {BRAND, HEADER_FILL}
TABLE_FONT_PT = 8.0

STYLE_SPECS = {
    "Normal": (BODY_FONT, 9.5, BODY_COLOR, None),
    "Heading 1": (BODY_FONT, 15, H1_COLOR, True),
    "Heading 2": (BODY_FONT, 11, H2_COLOR, True),
    "Heading 3": (BODY_FONT, 10, H3_COLOR, True),
    "Title": (BODY_FONT, 26, H1_COLOR, True),
    "Subtitle": (BODY_FONT, 11, H2_COLOR, False),
}


def _hex(c):
    return (c or "").lstrip("#").upper()


def configure_styles(doc):
    for name, (font, size, color, bold) in STYLE_SPECS.items():
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = font
        # ensure the explicit font wins over any theme font reference
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rf.set(qn(attr), font)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme"):
            if rf.get(qn(attr)) is not None:
                del rf.attrib[qn(attr)]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            st.font.bold = bold


def add_run(p, r, force_white=False, font=None, size=None):
    run = p.add_run(r.get("text", ""))
    if r.get("bold"):
        run.bold = True
    if r.get("italic"):
        run.italic = True
    if r.get("underline"):
        run.underline = True
    if size:
        run.font.size = Pt(size)
    elif r.get("size"):
        run.font.size = Pt(r["size"])
    if font:
        run.font.name = font
    if force_white:
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    elif r.get("color"):
        hexv = _hex(r["color"])
        if hexv == BRAND:
            hexv = H1_COLOR
        run.font.color.rgb = RGBColor.from_string(hexv)
    if r.get("break"):
        run.add_break()
    return run


def add_paragraph(doc, block):
    style = None
    lst = block.get("list")
    if lst:
        base = "List Number" if lst.get("type") == "numbered" else "List Bullet"
        lvl = lst.get("level", 0)
        style = f"{base} {lvl + 1}" if lvl else base
    try:
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    except KeyError:
        p = doc.add_paragraph()
    align = block.get("alignment")
    if align in ALIGN:
        p.alignment = ALIGN[align]
    if "runs" in block:
        for r in block["runs"]:
            add_run(p, r)
    elif "text" in block:
        add_run(p, {"text": block["text"]})
    return p


def shade_cell(cell, color_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def add_table(doc, block):
    rows = block["rows"]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            tc = table.cell(ri, ci)
            tc.text = ""
            p = tc.paragraphs[0]
            fill = _hex(cell.get("shading", "")) if cell.get("shading") else None
            if fill in DARK_FILLS:
                fill = HEADER_FILL
                white = True
            else:
                white = False
            add_run(
                p,
                {"text": cell.get("text", ""), "bold": white},
                force_white=white,
                font=BODY_FONT,
                size=TABLE_FONT_PT,
            )
            if fill:
                shade_cell(tc, fill)
    return table


def main(src):
    spec = json.loads(Path(src).read_text(encoding="utf-8"))
    doc = Document()
    configure_styles(doc)
    for section in spec.get("sections", []):
        for block in section.get("children", []):
            t = block.get("type")
            if t == "heading":
                doc.add_heading(block.get("text", ""), level=block.get("level", 1))
            elif t == "paragraph":
                add_paragraph(doc, block)
            elif t == "table":
                add_table(doc, block)
            elif t == "pageBreak":
                doc.add_page_break()
    out = Path(src).with_suffix("")  # drop .json -> ...docx
    doc.save(str(out))
    print(f"Saved: {out}")


if __name__ == "__main__":
    main(sys.argv[1])
